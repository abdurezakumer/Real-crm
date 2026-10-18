import csv
import json
import re
from collections.abc import Iterator
from dataclasses import dataclass
from itertools import islice
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from psycopg import Connection

from realestate_crm.phone import normalize_phone

ALIASES = {
    "source_row_number": {
        "no",
        "number",
        "id",
        "serial",
        "serialno",
        "sno",
        "srno",
        "index",
        "row",
    },
    "name": {
        "name",
        "fullname",
        "contactname",
        "customername",
        "clientname",
        "leadname",
        "contactperson",
    },
    "phone": {
        "phone",
        "phonenumber",
        "phoneno",
        "phonenum",
        "mobile",
        "mobilenumber",
        "mobileno",
        "cell",
        "cellphone",
        "telephone",
        "tel",
        "contactnumber",
        "contactno",
        "whatsapp",
        "whatsappnumber",
    },
}


class ColumnMappingError(ValueError):
    def __init__(self, headers: list[str], candidates: dict[str, list[int]]):
        self.headers = headers
        self.candidates = candidates
        super().__init__("Column selection is ambiguous or required columns are missing")


@dataclass(frozen=True)
class ImportPreview:
    headers: list[str]
    mapping: dict[str, int]
    rows: list[dict[str, str]]


def normalized_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def detect_mapping(headers: list[str], selected: dict[str, int] | None = None) -> dict[str, int]:
    if selected:
        mapping = {key: int(value) for key, value in selected.items()}
        if any(index < 0 or index >= len(headers) for index in mapping.values()):
            raise ColumnMappingError(headers, {})
        if "name" not in mapping or "phone" not in mapping:
            raise ColumnMappingError(headers, {})
        return mapping
    found: dict[str, list[int]] = {}
    for key, aliases in ALIASES.items():
        matched = [i for i, header in enumerate(headers) if normalized_header(header) in aliases]
        if matched:
            found[key] = matched
    if "phone" not in found:
        ambiguous_number_columns = [
            index
            for index, header in enumerate(headers)
            if normalized_header(header) in {"no", "number"}
        ]
        if ambiguous_number_columns:
            # Some lead sheets label the phone column simply "No.". When a
            # separate phone header exists, the same label stays a serial ID.
            found.pop("source_row_number", None)
            found["phone"] = ambiguous_number_columns
    if (
        "name" not in found
        or "phone" not in found
        or len(found.get("name", [])) != 1
        or len(found.get("phone", [])) != 1
    ):
        raise ColumnMappingError(headers, found)
    return {key: indices[0] for key, indices in found.items()}


def _rows(path: Path, extension: str, start_row: int = 1) -> Iterator[tuple[int, tuple[Any, ...]]]:
    if extension == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            for row_number, row in enumerate(csv.reader(handle), start=1):
                if row_number >= start_row:
                    yield row_number, tuple(row)
    elif extension == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            for row_number, row in enumerate(
                sheet.iter_rows(min_row=start_row, values_only=True), start=start_row
            ):
                yield row_number, tuple(row)
        finally:
            workbook.close()
    else:
        raise ValueError("Only .xlsx and .csv imports are supported")


def preview_file(
    path: Path, filename: str, selected: dict[str, int] | None = None
) -> ImportPreview:
    ext = Path(filename).suffix.lower()
    header_iter = _rows(path, ext, 1)
    try:
        _, header_values = next(header_iter)
    except StopIteration as exc:
        raise ValueError("The file has no header row") from exc
    headers = [str(value or "").strip() for value in header_values]
    mapping = detect_mapping(headers, selected)
    preview: list[dict[str, str]] = []
    for _, row in islice(_rows(path, ext, 2), 20):
        name = str(row[mapping["name"]] if mapping["name"] < len(row) else "").strip()
        phone_raw = row[mapping["phone"]] if mapping["phone"] < len(row) else ""
        phone_text = str(phone_raw or "").strip()
        phone = normalize_phone(phone_raw) or phone_text
        preview.append({"name": name[:100], "phone": phone[:50]})
    return ImportPreview(headers, mapping, preview)


def data_rows(
    path: Path, filename: str, mapping: dict[str, int], start_row: int, limit: int
) -> list[tuple[int, str, str, str | None]]:
    ext = Path(filename).suffix.lower()
    output = []
    max_index = max(mapping.values())
    for row_number, row in islice(_rows(path, ext, start_row), limit):
        if not any(str(value or "").strip() for value in row):
            output.append((row_number, "", "", None))
            continue
        if max_index >= len(row):
            output.append((row_number, "", "", None))
            continue
        name = str(row[mapping["name"]] or "").strip()
        phone_raw = str(row[mapping["phone"]] or "").strip()
        source_index = mapping.get("source_row_number")
        source_number = (
            str(row[source_index]).strip()
            if source_index is not None and source_index < len(row)
            else None
        )
        output.append((row_number, name, phone_raw, source_number))
    return output


def process_import_batch(
    conn: Connection, job: dict[str, Any], path: Path, batch_size: int, max_rows: int
) -> dict[str, Any]:
    mapping = job["header_mapping"]
    if isinstance(mapping, str):
        mapping = json.loads(mapping)
    scanned = data_rows(
        path, job["original_filename"], mapping, job["next_row_number"], batch_size + 1
    )
    too_many = any(row[0] > max_rows + 1 for row in scanned)
    rows = [row for row in scanned if row[0] <= max_rows + 1][:batch_size]
    errors = list(job["error_summary"] or [])
    accepted: list[tuple[int, str, str, str]] = []
    duplicate_rows = 0
    invalid_rows = 0
    local_seen: set[str] = set()
    for row_number, name, raw_phone, source_number in rows:
        if not name and not raw_phone and source_number is None:
            continue
        normalized = normalize_phone(raw_phone)
        reason = None
        if not name:
            reason = "Missing contact name"
        elif not normalized:
            reason = "Invalid or missing phone number"
        elif normalized in local_seen:
            duplicate_rows += 1
            reason = "Duplicate phone number in this import"
        if reason:
            if reason != "Duplicate phone number in this import":
                invalid_rows += 1
            if len(errors) < 200:
                errors.append({"row": row_number, "reason": reason})
            continue
        local_seen.add(normalized)
        source_row = source_number or str(row_number)
        accepted.append((row_number, name, normalized, source_row[:50]))

    phone_list = [item[2] for item in accepted]
    existing = set()
    if phone_list:
        existing = {
            row["normalized_phone"]
            for row in conn.execute(
                "select normalized_phone from contacts where normalized_phone = any(%s)",
                (phone_list,),
            ).fetchall()
        }
    policy = job["import_policy"]
    already = existing
    if policy == "SKIP_EXISTING":
        skipped = sum(1 for item in accepted if item[2] in already)
        duplicate_rows += skipped
        accepted_to_write = [item for item in accepted if item[2] not in already]
    else:
        skipped = 0
        duplicate_rows += len(already)
        accepted_to_write = accepted

    if accepted_to_write:
        lead_source = job.get("lead_source")
        if job.get("campaign_id"):
            campaign = conn.execute(
                "select source from campaigns where id=%s", (job["campaign_id"],)
            ).fetchone()
            lead_source = lead_source or (campaign["source"] if campaign else None)
        if policy == "UPDATE_EXISTING":
            statement = """insert into contacts (name, phone_number, normalized_phone, source_row_number,
                  assigned_telegram_user_id,assigned_at,assignment_source,import_batch_id,lead_source)
                values (%s,%s,%s,%s,%s,now(),'IMPORT',%s,%s) on conflict (normalized_phone) where normalized_phone is not null
                do update set name=excluded.name, phone_number=excluded.phone_number,
                  source_row_number=excluded.source_row_number,import_batch_id=excluded.import_batch_id,
                  lead_source=coalesce(excluded.lead_source,contacts.lead_source),updated_at=now()"""
        else:
            statement = """insert into contacts (name, phone_number, normalized_phone, source_row_number,
                  assigned_telegram_user_id,assigned_at,assignment_source,import_batch_id,lead_source)
                values (%s,%s,%s,%s,%s,now(),'IMPORT',%s,%s) on conflict (normalized_phone) where normalized_phone is not null do nothing"""
        conn.executemany(
            statement,
            [
                (
                    name,
                    phone,
                    phone,
                    int(source_row) if source_row.isdigit() else row_number,
                    job.get("telegram_user_id"),
                    job["id"],
                    lead_source,
                )
                for row_number, name, phone, source_row in accepted_to_write
            ],
        )
    if accepted:
        ids = conn.execute(
            "select id, normalized_phone from contacts where normalized_phone = any(%s)",
            (phone_list,),
        ).fetchall()
        selected_ids = [
            row["id"] for row in ids if row["normalized_phone"] in {item[2] for item in accepted}
        ]
        if job.get("telegram_user_id"):
            conn.execute(
                """insert into lead_assignment_history(contact_id,new_owner_telegram_user_id,
                     actor_telegram_user_id,assignment_source,reason,idempotency_key)
                   select c.id,c.assigned_telegram_user_id,%s,'IMPORT','Assigned on first import',
                     'import:'||c.id::text from contacts c
                   where c.id=any(%s) and c.assignment_source='IMPORT'
                     and c.assigned_telegram_user_id=%s and not exists(
                       select 1 from lead_assignment_history h where h.contact_id=c.id
                         and h.assignment_source='IMPORT')
                   on conflict(idempotency_key) do nothing""",
                (job["telegram_user_id"], selected_ids, job["telegram_user_id"]),
            )
        if job["campaign_id"]:
            conn.executemany(
                "insert into campaign_contacts (campaign_id, contact_id) values (%s,%s) on conflict do nothing",
                [(job["campaign_id"], contact_id) for contact_id in selected_ids],
            )
    imported = len(accepted_to_write) if policy != "SKIP_EXISTING" else len(accepted_to_write)
    if policy == "ADD_TO_CAMPAIGN" and job["campaign_id"]:
        imported = len(accepted)
    next_row = rows[-1][0] + 1 if rows else job["next_row_number"]
    if too_many and len(errors) < 200:
        errors.append(
            {
                "row": max_rows + 2,
                "reason": f"Maximum row limit ({max_rows}) exceeded; later rows were not imported",
            }
        )
    done = too_many or len(scanned) <= batch_size
    status = "FAILED" if too_many else ("COMPLETED" if done else "PROCESSING")
    total = max(0, next_row - 2) if done else job["total_rows"]
    conn.execute(
        """update import_jobs set status=%s, next_row_number=%s, total_rows=coalesce(%s,total_rows),
             imported_rows=imported_rows+%s, skipped_rows=skipped_rows+%s,
             duplicate_rows=duplicate_rows+%s, invalid_rows=invalid_rows+%s,
             error_summary=%s, completed_at=case when %s in ('COMPLETED','FAILED') then now() else null end,
             updated_at=now() where id=%s""",
        (
            status,
            next_row,
            total,
            imported,
            skipped,
            duplicate_rows,
            invalid_rows,
            json.dumps(errors),
            status,
            job["id"],
        ),
    )
    return {
        "status": status,
        "processed": len(rows),
        "imported": imported,
        "skipped": skipped,
        "duplicates": duplicate_rows,
        "invalid": invalid_rows,
        "next_row_number": next_row,
    }
