"""Streaming Excel snapshots of all application data in the public schema."""

import base64
import json
import re
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any
from uuid import UUID

from openpyxl import Workbook
from psycopg import sql
from psycopg.rows import tuple_row

from realestate_crm.database import connection

EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_DATA_ROWS = EXCEL_MAX_ROWS - 1
EXCEL_MAX_CELL_CHARS = 30_000
FETCH_BATCH_SIZE = 1_000


def _sheet_name(base: str, used: set[str], suffix: str = "") -> str:
    cleaned = re.sub(r"[\\/*?:\[\]]", "_", base).strip(" '") or "data"
    stem = cleaned[: 31 - len(suffix)]
    candidate = f"{stem}{suffix}"
    number = 2
    while candidate.casefold() in used:
        numbered_suffix = f"_{number}"
        candidate = f"{cleaned[: 31 - len(numbered_suffix)]}{numbered_suffix}"
        number += 1
    used.add(candidate.casefold())
    return candidate


def _cell_value(value: Any) -> str | int | float | bool | None:
    if value is None or isinstance(value, (bool, int)):
        if isinstance(value, int) and not isinstance(value, bool) and len(str(abs(value))) > 15:
            return str(value)
        return value
    if isinstance(value, str):
        text = value
    elif isinstance(value, (datetime, date, Decimal, UUID)):
        text = value.isoformat() if isinstance(value, (datetime, date)) else str(value)
    elif isinstance(value, float):
        text = repr(value)
    elif isinstance(value, bytes | bytearray | memoryview):
        text = base64.b64encode(bytes(value)).decode("ascii")
    elif isinstance(value, (dict, list, tuple)):
        text = json.dumps(value, ensure_ascii=False, default=str, separators=(",", ":"))
    else:
        text = str(value)
    if text.lstrip(" \t\r\n")[:1] in {"=", "+", "-", "@"}:
        text = "'" + text
    return text


def build_database_backup(path: Path) -> dict[str, Any]:
    """Write a consistent, streaming XLSX snapshot of every public base table."""
    workbook = Workbook(write_only=True)
    used_sheets = {"backup manifest"}
    manifest = workbook.create_sheet("Backup Manifest")
    manifest.append(["Created (UTC)", datetime.now(UTC).isoformat()])
    manifest.append(["Scope", "All base tables in the PostgreSQL public schema"])
    manifest.append([])
    manifest.append(["Table", "Rows", "Data sheets", "Overflow sheets"])
    table_count = 0
    total_rows = 0

    with connection() as conn:
        conn.execute("set transaction isolation level repeatable read read only")
        tables = conn.execute(
            """select table_name from information_schema.tables
               where table_schema='public' and table_type='BASE TABLE'
               order by table_name"""
        ).fetchall()

        for table_number, table_row in enumerate(tables, start=1):
            table_name = table_row["table_name"]
            overflow_sheet_count = 0
            row_count = 0
            data_sheet_count = 1
            data_sheet = workbook.create_sheet(_sheet_name(table_name, used_sheets))
            data_rows_on_sheet = 0
            overflow_sheet = None
            overflow_rows_on_sheet = EXCEL_MAX_DATA_ROWS

            with conn.cursor(
                name=f"crm_backup_{table_number}", row_factory=tuple_row
            ) as cursor:
                cursor.execute(
                    sql.SQL("select * from {}.{}").format(
                        sql.Identifier("public"), sql.Identifier(table_name)
                    )
                )
                columns = [column.name for column in cursor.description]
                data_sheet.append(columns)
                while batch := cursor.fetchmany(FETCH_BATCH_SIZE):
                    for row in batch:
                        if data_rows_on_sheet >= EXCEL_MAX_DATA_ROWS:
                            data_sheet_count += 1
                            data_sheet = workbook.create_sheet(
                                _sheet_name(table_name, used_sheets, f"_{data_sheet_count}")
                            )
                            data_sheet.append(columns)
                            data_rows_on_sheet = 0

                        row_count += 1
                        total_rows += 1
                        output_row = []
                        overflow_values: list[tuple[str, str]] = []
                        for column_name, value in zip(columns, row, strict=True):
                            cell = _cell_value(value)
                            if isinstance(cell, str) and len(cell) > EXCEL_MAX_CELL_CHARS:
                                overflow_values.append(
                                    (column_name, cell[EXCEL_MAX_CELL_CHARS:])
                                )
                                cell = cell[:EXCEL_MAX_CELL_CHARS]
                            output_row.append(cell)
                        data_sheet.append(output_row)
                        data_rows_on_sheet += 1

                        for column_name, remainder in overflow_values:
                            for part_number, offset in enumerate(
                                range(0, len(remainder), EXCEL_MAX_CELL_CHARS), start=2
                            ):
                                if (
                                    overflow_sheet is None
                                    or overflow_rows_on_sheet >= EXCEL_MAX_DATA_ROWS
                                ):
                                    overflow_sheet_count += 1
                                    overflow_sheet = workbook.create_sheet(
                                        _sheet_name(
                                            f"{table_name}_long_text",
                                            used_sheets,
                                            f"_{overflow_sheet_count}",
                                        )
                                    )
                                    overflow_sheet.append(
                                        ["Source row", "Column", "Part", "Text"]
                                    )
                                    overflow_rows_on_sheet = 0
                                overflow_sheet.append(
                                    [
                                        row_count,
                                        column_name,
                                        part_number,
                                        remainder[offset : offset + EXCEL_MAX_CELL_CHARS],
                                    ]
                                )
                                overflow_rows_on_sheet += 1

            manifest.append(
                [table_name, row_count, data_sheet_count, overflow_sheet_count]
            )
            table_count += 1

    workbook.save(path)
    return {
        "table_count": table_count,
        "row_count": total_rows,
        "file_size_bytes": path.stat().st_size,
    }
